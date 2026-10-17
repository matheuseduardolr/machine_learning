"""Leitura e gravação da planilha de condomínios."""

from __future__ import annotations

from copy import copy

from openpyxl.styles import Alignment, Font, PatternFill

from project_clustering_kmeans.settings import (
    COL_BAIRRO,
    COL_CIDADE,
    COL_COR,
    COL_ENDERECO,
    COL_LAT,
    COL_LON,
    COL_NOME,
    COL_NUCLEO,
    COL_ORDEM,
    HEADER_ROW,
    NUCLEUS_COLORS,
)


def read_rows(ws) -> list[dict]:
    rows = []
    for index in range(HEADER_ROW + 1, ws.max_row + 1):
        ordem = ws.cell(index, COL_ORDEM).value
        nome = ws.cell(index, COL_NOME).value
        if ordem is None or nome is None:
            continue
        rows.append(
            {
                "row": index,
                "ordem": int(ordem),
                "nome": str(nome).strip(),
                "endereco": str(ws.cell(index, COL_ENDERECO).value or "").strip(),
                "bairro": str(ws.cell(index, COL_BAIRRO).value or "").strip(),
                "cidade": str(ws.cell(index, COL_CIDADE).value or "").strip(),
            }
        )
    return rows


def coordinates_from_sheet(ws, rows: list[dict]) -> dict[int, dict]:
    located: dict[int, dict] = {}
    for item in rows:
        latitude = ws.cell(item["row"], COL_LAT).value
        longitude = ws.cell(item["row"], COL_LON).value
        if latitude is None or longitude is None:
            continue
        located[item["ordem"]] = {"lat": float(latitude), "lon": float(longitude)}
    return located


def tint(hex_color: str, white_ratio: float = 0.84) -> str:
    channels = [int(hex_color[index : index + 2], 16) for index in (0, 2, 4)]
    mixed = [round(channel + (255 - channel) * white_ratio) for channel in channels]
    return "".join(f"{channel:02X}" for channel in mixed)


def write_results(ws, rows: list[dict], located: dict[int, dict], nuclei: dict[int, int]) -> None:
    merged = [str(item) for item in ws.merged_cells.ranges]
    for ref in merged:
        if ref.startswith("A1:") or ref.startswith("A2:"):
            ws.unmerge_cells(ref)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=COL_COR)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=COL_COR)

    header_src = ws.cell(HEADER_ROW, COL_ORDEM)
    headers = {
        COL_LAT: "Latitude",
        COL_LON: "Longitude",
        COL_NUCLEO: "Núcleo",
        COL_COR: "Cor",
    }
    for column, title in headers.items():
        cell = ws.cell(HEADER_ROW, column, title)
        cell.font = copy(header_src.font)
        cell.fill = copy(header_src.fill)
        cell.alignment = copy(header_src.alignment)
        cell.border = copy(header_src.border)

    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 12
    ws.column_dimensions["J"].width = 14

    for item in rows:
        result = located.get(item["ordem"])
        lat_cell = ws.cell(item["row"], COL_LAT)
        lon_cell = ws.cell(item["row"], COL_LON)
        nucleus_cell = ws.cell(item["row"], COL_NUCLEO)
        color_cell = ws.cell(item["row"], COL_COR)
        if result is None:
            lat_cell.value = None
            lon_cell.value = None
            nucleus_cell.value = None
            color_cell.value = None
            continue

        nucleus = nuclei[item["ordem"]]
        color_name, color_hex = NUCLEUS_COLORS[nucleus - 1]
        solid = PatternFill(fill_type="solid", fgColor=color_hex)
        light = PatternFill(fill_type="solid", fgColor=tint(color_hex))
        font_color = "000000" if color_name == "Amarelo" else "FFFFFF"

        lat_cell.value = round(result["lat"], 6)
        lon_cell.value = round(result["lon"], 6)
        lat_cell.number_format = "0.000000"
        lon_cell.number_format = "0.000000"
        nucleus_cell.value = nucleus
        color_cell.value = color_name
        for column in range(COL_ORDEM, COL_CIDADE + 1):
            ws.cell(item["row"], column).fill = light
        lat_cell.fill = light
        lon_cell.fill = light
        for cell in (nucleus_cell, color_cell):
            cell.fill = solid
            cell.font = Font(color=font_color, bold=True)
            cell.alignment = Alignment(horizontal="center")
